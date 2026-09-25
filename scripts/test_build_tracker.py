"""Tests for the Summary sheet written by build_tracker.py (run: python -m unittest scripts.test_build_tracker)."""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_tracker  # noqa: E402

DATA = {"milestones": [{"name": "M1"}, {"name": "M2"}]}
# Only columns A (milestone) and E (status) matter to the Summary.
TASK_ROWS = [
    ["M1", "", "", "", "Completed"],
    ["M1", "", "", "", "Completed"],
    ["M1", "", "", "", "Not Started"],
    ["M2", "", "", "", "In Progress"],
    ["M2", "", "", "", "Pending Verification"],
]
TRACE_ROWS = [
    ["FR-01.1", "", "", "", "Completed"],
    ["FR-01.2", "", "", "", "Partial"],
    ["FR-01.3", "", "", "", "Partial"],
    ["FR-01.4", "", "", "", "Not Started"],
]


def read_summary(data_only):
    wb = Workbook()
    build_tracker.build_summary(wb, DATA, TASK_ROWS, TRACE_ROWS)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "summary.xlsx"
        wb.save(path)
        return load_workbook(path, data_only=data_only)["Summary"]


class SummarySheetTest(unittest.TestCase):
    def test_no_numeric_cell_reads_as_none_without_recalculation(self):
        sheet = read_summary(data_only=True)
        for row in sheet.iter_rows(min_row=2, max_row=4, min_col=2, max_col=7):
            for cell in row:
                self.assertIsNotNone(cell.value, f"{cell.coordinate} is empty in data_only mode")

    def test_counts_match_the_task_rows(self):
        sheet = read_summary(data_only=True)
        # Columns: Not Started, In Progress, Pending Verification, Completed, Total, % Completed.
        self.assertEqual([c.value for c in sheet[2][1:7]], [1, 0, 0, 2, 3, 2 / 3])
        self.assertEqual([c.value for c in sheet[3][1:7]], [0, 1, 1, 0, 2, 0])
        self.assertEqual([c.value for c in sheet[4][1:7]], [1, 1, 1, 2, 5, 0.4])

    def test_requirement_counts_match_the_traceability_rows(self):
        sheet = read_summary(data_only=True)
        self.assertEqual([c.value for c in sheet[6][1:4]], ["Not Started", "Partial", "Completed"])
        self.assertEqual([c.value for c in sheet[7][1:5]], [1, 2, 1, 4])

    def test_cells_hold_plain_numbers_not_formulas(self):
        # openpyxl cannot store a formula together with its cached result, so the Summary stores numbers.
        sheet = read_summary(data_only=False)
        self.assertEqual(sheet["E2"].value, 2)


if __name__ == "__main__":
    unittest.main()
