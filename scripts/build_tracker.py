#!/usr/bin/env python3
"""Builds the SELA progress tracker workbook from docs/tracking/tracker.json.

Outputs (in docs/tracking/): SELA_Progress_Tracker.xlsx plus tracker.csv, changelog.csv and
traceability.csv for plain-text diffs and Excel import.

Requirement descriptions are parsed straight from docs/FSD_Sela.md so the Traceability sheet never
drifts from the spec, and every commit hash in the data is verified against git.

Usage:
    python scripts/build_tracker.py           validate, then write the workbook and CSVs
    python scripts/build_tracker.py --check   validate only (no files written)
"""
import csv
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TRACK_DIR = ROOT / "docs" / "tracking"
DATA_FILE = TRACK_DIR / "tracker.json"
FSD_FILE = ROOT / "docs" / "FSD_Sela.md"
XLSX_FILE = TRACK_DIR / "SELA_Progress_Tracker.xlsx"

TASK_STATUSES = ["Not Started", "In Progress", "Pending Verification", "Completed"]
REQ_STATUSES = ["Not Started", "Partial", "Completed"]

STATUS_FILL = {
    "Not Started": "E7E6E6",
    "In Progress": "FFF2CC",
    "Pending Verification": "FCE4D6",
    "Completed": "C6E0B4",
    "Partial": "FFF2CC",
}

HASH_RE = re.compile(r"\b[0-9a-f]{7,40}\b")
TASK_REF_RE = re.compile(r"\b(?:T\d+\.\d+|M\d\.\d+|BL\.\d+)\b")
FR_ID_RE = re.compile(r"\bFR-(?:SEC\.\d+|\d{2}\.\d+)\b")
FSD_ROW_RE = re.compile(r"^\|\s*(FR-(?:SEC\.\d+|\d{2}\.\d+))\s*\|\s*(.*?)\s*\|\s*$", re.MULTILINE)


def load_data():
    with DATA_FILE.open(encoding="utf-8") as fh:
        return json.load(fh)


def parse_fsd_requirements():
    """Returns {FR-ID: description}, keeping the first (specification) occurrence of each ID."""
    text = FSD_FILE.read_text(encoding="utf-8")
    found = {}
    for fr_id, description in FSD_ROW_RE.findall(text):
        found.setdefault(fr_id, description)
    return found


def git_has_commit(commit):
    result = subprocess.run(
        ["git", "cat-file", "-e", f"{commit}^{{commit}}"],
        cwd=ROOT, capture_output=True, check=False,
    )
    return result.returncode == 0


def validate(data, fsd_requirements):
    errors = []
    milestone_ids = [m["id"] for m in data["milestones"]]
    task_ids = [t["id"] for t in data["tasks"]]

    for dup in {i for i in task_ids if task_ids.count(i) > 1}:
        errors.append(f"duplicate task id {dup}")

    for task in data["tasks"]:
        tid = task["id"]
        if task["milestone"] not in milestone_ids:
            errors.append(f"{tid}: unknown milestone {task['milestone']}")
        if task["status"] not in TASK_STATUSES:
            errors.append(f"{tid}: invalid status {task['status']!r}")
        if task["status"] == "Completed" and not task["green"]:
            errors.append(f"{tid}: a Completed task needs a GREEN commit (or an explicit n/a reason)")
        for field in ("red", "green"):
            for commit in HASH_RE.findall(task[field]):
                if not git_has_commit(commit):
                    errors.append(f"{tid}: {field} commit {commit} not found in git")
        for fr_id in FR_ID_RE.findall(task["frs"]):
            if fr_id not in fsd_requirements:
                errors.append(f"{tid}: {fr_id} is not a requirement in the FSD")

    for entry in data["changelog"]:
        if entry["task"] not in task_ids:
            errors.append(f"changelog entry references unknown task {entry['task']}")
        for commit in HASH_RE.findall(entry["commits"]):
            if not git_has_commit(commit):
                errors.append(f"changelog {entry['task']}: commit {commit} not found in git")

    for fr_id, req in data["requirements"].items():
        if fr_id not in fsd_requirements:
            errors.append(f"requirements: {fr_id} is not in the FSD")
        if req.get("status", "Not Started") not in REQ_STATUSES:
            errors.append(f"requirements {fr_id}: invalid status {req.get('status')!r}")
        for ref in TASK_REF_RE.findall(req.get("tasks", "")):
            if ref not in task_ids:
                errors.append(f"requirements {fr_id}: unknown task {ref}")

    completed = {t["id"] for t in data["tasks"] if t["status"] == "Completed"}
    logged = {e["task"] for e in data["changelog"]}
    for missing in sorted(completed - logged):
        errors.append(f"{missing}: Completed but has no change log entry")

    return errors


def _fr_sort_key(fr_id):
    prefix, number = fr_id.split(".")
    return (prefix == "FR-SEC", prefix, int(number))


def _task_sort_key(task_id):
    head, _, tail = task_id.partition(".")
    return (head, int(tail) if tail.isdigit() else 0)


def milestone_name(data, milestone_id):
    for m in data["milestones"]:
        if m["id"] == milestone_id:
            return m["name"]
    return milestone_id


def requirement_rows(data, fsd_requirements):
    prefix_milestone = data["requirement_defaults"]
    prefix_section = data["fsd_sections"]
    rows = []
    for fr_id in sorted(fsd_requirements, key=_fr_sort_key):
        prefix = fr_id.split(".")[0]
        override = data["requirements"].get(fr_id, {})
        rows.append({
            "id": fr_id,
            "description": fsd_requirements[fr_id],
            "section": prefix_section.get(prefix, ""),
            "milestone": override.get("milestone") or prefix_milestone.get(prefix, ""),
            "status": override.get("status", "Not Started"),
            "tasks": override.get("tasks", ""),
            "evidence": override.get("evidence", ""),
        })
    return rows


def tracker_rows(data):
    order = {m["id"]: i for i, m in enumerate(data["milestones"])}
    tasks = sorted(data["tasks"], key=lambda t: (order[t["milestone"]], _task_sort_key(t["id"])))
    return [
        [
            milestone_name(data, t["milestone"]), t["id"], f"{t['task']} — {t['description']}", t["module"],
            t["status"], t["red"], t["green"], t["coverage"], t["frs"], t["deliverables"], t["updated"],
        ]
        for t in tasks
    ]


def changelog_rows(data):
    order = {m["id"]: i for i, m in enumerate(data["milestones"])}
    entries = sorted(data["changelog"], key=lambda e: (order[e["milestone"]], e["module"], _task_sort_key(e["task"])))
    return [
        [milestone_name(data, e["milestone"]), e["module"], e["task"], e["summary"], e["breaking"],
         e["decisions"], e["commits"], e["date"]]
        for e in entries
    ]


def trace_rows(data, requirements):
    return [
        [r["id"], r["description"], r["section"], milestone_name(data, r["milestone"]), r["status"], r["tasks"], r["evidence"]]
        for r in requirements
    ]


TRACKER_HEADERS = ["Phase / Milestone", "Task ID", "Task & Description", "Module / Service", "Status",
                   "TDD: RED commit", "TDD: GREEN commit", "Test coverage", "FR IDs", "Key Deliverables / Artifacts",
                   "Last Updated"]
CHANGELOG_HEADERS = ["Phase / Milestone", "Module / Service", "Task ID", "Summary of implemented functionality",
                     "Breaking changes / structural refactors", "Deviations & technical decisions", "Commits", "Date"]
TRACE_HEADERS = ["FR ID", "Requirement (from FSD)", "FSD section", "Milestone", "Status", "Tracker task(s)", "Evidence"]

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_table(ws, headers, rows, widths, freeze="C2"):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def add_status_formatting(ws, column_letter, last_row):
    for status, color in STATUS_FILL.items():
        ws.conditional_formatting.add(
            f"{column_letter}2:{column_letter}{last_row}",
            FormulaRule(formula=[f'${column_letter}2="{status}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)),
        )


def add_status_dropdown(ws, column_letter, last_row, statuses):
    dv = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{last_row}")


def _share(part, whole):
    return part / whole if whole else 0


def build_summary(wb, data, task_rows, trace_rows_):
    """Writes computed numbers, not formulas: openpyxl cannot store a formula together with its cached
    result, so formula cells read as None for pandas/openpyxl (data_only=True). The workbook is
    regenerated from tracker.json on every run, so there is nothing for a formula to keep live."""
    summary = wb.create_sheet("Summary", 1)
    summary.append(["Phase / Milestone"] + TASK_STATUSES + ["Total", "% Completed"])
    first = 2
    totals = [0] * len(TASK_STATUSES)
    for milestone in data["milestones"]:
        counts = [
            sum(1 for row in task_rows if row[0] == milestone["name"] and row[4] == status)
            for status in TASK_STATUSES
        ]
        totals = [t + c for t, c in zip(totals, counts)]
        summary.append([milestone["name"]] + counts + [sum(counts), _share(counts[-1], sum(counts))])
    total_row = first + len(data["milestones"])
    summary.append(["All tasks"] + totals + [sum(totals), _share(totals[-1], sum(totals))])
    for r in range(first, total_row + 1):
        summary.cell(row=r, column=7).number_format = "0%"
    summary.append([])
    summary.append(["Requirements (FR IDs)"] + REQ_STATUSES + ["Total"])
    req_header = summary.max_row
    req_counts = [sum(1 for row in trace_rows_ if row[4] == status) for status in REQ_STATUSES]
    summary.append(["All FR IDs"] + req_counts + [sum(req_counts)])
    for row in (1, req_header):
        for cell in summary[row]:
            if cell.value:
                cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
    summary.column_dimensions["A"].width = 46
    for col in "BCDEFG":
        summary.column_dimensions[col].width = 20
    summary.freeze_panes = "B2"


def build_guide(wb, data):
    guide = wb.active
    guide.title = "Guide"
    lines = [
        ("SELA progress tracker", True),
        (f"Data as of {data['as_of']}. Generated from docs/tracking/tracker.json and docs/FSD_Sela.md; do not edit this workbook by hand.", False),
        ("To update: edit tracker.json, then run `python scripts/build_tracker.py` (see docs/tracking/README.md).", False),
        ("", False),
        ("Task statuses", True),
        ("Not Started: nothing committed for the task.", False),
        ("In Progress: RED test committed or work under way.", False),
        ("Pending Verification: implemented, but a stated verification step is still open (for example against real infrastructure).", False),
        ("Completed: RED and GREEN commits recorded, coverage >= 80% where applicable, scripts/check.sh passes, docs synced.", False),
        ("", False),
        ("Requirement statuses (Traceability sheet): Not Started, Partial, Completed.", False),
        ("Sheets: Summary (counts computed at build time), Tracker (one row per task and module), Change Log, Traceability (FR IDs).", False),
    ]
    for index, (text, bold) in enumerate(lines, start=1):
        guide.cell(row=index, column=1, value=text).font = Font(bold=bold, size=14 if index == 1 else 11)
    guide.column_dimensions["A"].width = 130


def build_workbook(data, fsd_requirements):
    wb = Workbook()
    build_guide(wb, data)

    t_rows = tracker_rows(data)
    tracker = wb.create_sheet("Tracker")
    write_table(tracker, TRACKER_HEADERS, t_rows, [30, 9, 55, 34, 20, 18, 20, 22, 26, 70, 14])
    add_status_dropdown(tracker, "E", tracker.max_row, TASK_STATUSES)
    add_status_formatting(tracker, "E", tracker.max_row)

    c_rows = changelog_rows(data)
    changelog = wb.create_sheet("Change Log")
    write_table(changelog, CHANGELOG_HEADERS, c_rows, [30, 30, 9, 60, 55, 70, 30, 12])

    requirements = requirement_rows(data, fsd_requirements)
    r_rows = trace_rows(data, requirements)
    trace = wb.create_sheet("Traceability")
    write_table(trace, TRACE_HEADERS, r_rows, [11, 80, 11, 34, 14, 26, 60], freeze="B2")
    add_status_formatting(trace, "E", trace.max_row)
    add_status_dropdown(trace, "E", trace.max_row, REQ_STATUSES)

    build_summary(wb, data, t_rows, r_rows)
    return wb, t_rows, c_rows, r_rows


def write_csv(path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)


def main(argv):
    check_only = "--check" in argv
    data = load_data()
    fsd_requirements = parse_fsd_requirements()
    if not fsd_requirements:
        print("no FR rows found in docs/FSD_Sela.md", file=sys.stderr)
        return 1

    errors = validate(data, fsd_requirements)
    if errors:
        print("tracker validation failed:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 1

    wb, t_rows, c_rows, r_rows = build_workbook(data, fsd_requirements)
    counts = {status: sum(1 for row in t_rows if row[4] == status) for status in TASK_STATUSES}
    print(f"tracker ok: {len(t_rows)} tasks {counts}; {len(c_rows)} change log entries; {len(r_rows)} requirements")
    if check_only:
        return 0

    wb.save(XLSX_FILE)
    write_csv(TRACK_DIR / "tracker.csv", TRACKER_HEADERS, t_rows)
    write_csv(TRACK_DIR / "changelog.csv", CHANGELOG_HEADERS, c_rows)
    write_csv(TRACK_DIR / "traceability.csv", TRACE_HEADERS, r_rows)
    print(f"wrote {XLSX_FILE.relative_to(ROOT)} and 3 CSV files")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
